"""
Excel-specific data models for structured spreadsheet parsing.

These models represent Excel/CSV-specific structures:
- TableRange: Detected table boundaries
- SheetInfo: Sheet metadata and structure
- CellInfo: Individual cell information
"""
from dataclasses import dataclass, field
from typing import List, Dict, Any, Optional, Tuple


@dataclass
class CellInfo:
    """
    Information about a single Excel cell.

    Captures cell value, type, formatting, and position.
    """
    row: int
    column: int
    value: Any
    data_type: str  # "text", "number", "date", "formula", "boolean", "empty"
    formula: Optional[str] = None
    is_merged: bool = False
    merge_range: Optional[Tuple[int, int, int, int]] = None  # (min_row, min_col, max_row, max_col)
    style_info: Dict[str, Any] = field(default_factory=dict)

    @property
    def coordinate(self) -> str:
        """Get Excel-style coordinate (e.g., 'A1')."""
        from openpyxl.utils import get_column_letter
        return f"{get_column_letter(self.column)}{self.row}"

    @property
    def is_header_likely(self) -> bool:
        """Heuristic: Is this cell likely a header?"""
        if self.style_info.get('bold', False):
            return True
        if self.data_type == "text" and self.row == 1:
            return True
        return False


@dataclass
class TableRange:
    """
    Detected table boundary within a sheet.

    Represents a contiguous data region with headers and data rows.
    """
    sheet_name: str
    start_row: int
    start_col: int
    end_row: int
    end_col: int
    header_row: Optional[int] = None
    has_headers: bool = True
    table_name: Optional[str] = None

    @property
    def num_rows(self) -> int:
        """Total rows in table (including header)."""
        return self.end_row - self.start_row + 1

    @property
    def num_cols(self) -> int:
        """Total columns in table."""
        return self.end_col - self.start_col + 1

    @property
    def data_rows(self) -> int:
        """Number of data rows (excluding header)."""
        if self.has_headers:
            return self.num_rows - 1
        return self.num_rows

    @property
    def range_str(self) -> str:
        """Excel-style range string (e.g., 'A1:D10')."""
        from openpyxl.utils import get_column_letter
        start_col_letter = get_column_letter(self.start_col)
        end_col_letter = get_column_letter(self.end_col)
        return f"{start_col_letter}{self.start_row}:{end_col_letter}{self.end_row}"

    def contains_cell(self, row: int, col: int) -> bool:
        """Check if cell is within this table range."""
        return (self.start_row <= row <= self.end_row and
                self.start_col <= col <= self.end_col)


@dataclass
class SheetInfo:
    """
    Metadata and structure information for an Excel sheet.

    Contains sheet-level information and detected tables.
    """
    name: str
    index: int
    max_row: int
    max_column: int
    has_data: bool
    tables: List[TableRange] = field(default_factory=list)
    merged_cells: List[Tuple[int, int, int, int]] = field(default_factory=list)
    has_formulas: bool = False
    has_merged_cells: bool = False
    column_types: Dict[int, str] = field(default_factory=dict)  # col_idx -> data_type

    @property
    def total_cells(self) -> int:
        """Approximate total cells with data."""
        return self.max_row * self.max_column

    @property
    def num_tables(self) -> int:
        """Number of detected tables in sheet."""
        return len(self.tables)

    @property
    def is_empty(self) -> bool:
        """Check if sheet is effectively empty."""
        return not self.has_data or (self.max_row == 0 and self.max_column == 0)

    def get_table_at_cell(self, row: int, col: int) -> Optional[TableRange]:
        """Find table containing given cell."""
        for table in self.tables:
            if table.contains_cell(row, col):
                return table
        return None


@dataclass
class ExcelMetadata:
    """
    Complete metadata for an Excel workbook.

    Contains workbook-level information and all sheets.
    """
    filename: str
    total_sheets: int
    sheets: List[SheetInfo] = field(default_factory=list)
    author: Optional[str] = None
    created: Optional[str] = None
    modified: Optional[str] = None

    @property
    def active_sheets(self) -> List[SheetInfo]:
        """Get sheets that contain data."""
        return [sheet for sheet in self.sheets if not sheet.is_empty]

    @property
    def total_tables(self) -> int:
        """Total number of tables across all sheets."""
        return sum(sheet.num_tables for sheet in self.sheets)

    @property
    def has_multiple_sheets(self) -> bool:
        """Check if workbook has multiple sheets with data."""
        return len(self.active_sheets) > 1


@dataclass
class CSVMetadata:
    """
    Metadata for CSV files.

    Contains CSV-specific information like delimiter, encoding, etc.
    """
    filename: str
    delimiter: str = ","
    encoding: str = "utf-8"
    has_header: bool = True
    num_rows: int = 0
    num_columns: int = 0
    column_names: List[str] = field(default_factory=list)
    column_types: Dict[str, str] = field(default_factory=dict)

    @property
    def is_tsv(self) -> bool:
        """Check if file is tab-separated."""
        return self.delimiter == "\t"

    @property
    def is_semicolon(self) -> bool:
        """Check if file uses semicolon separator."""
        return self.delimiter == ";"
