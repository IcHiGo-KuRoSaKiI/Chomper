"""
PPTX extractor: slides to clean Markdown, in reading order.

Output decisions:

- **Slide header.** Each slide starts with ``## Slide N: <title>`` (``## Slide
  N`` when it has no title). The title is not repeated below the header.
  Hidden slides are marked ``(hidden)`` and listed in ``metadata``.
- **Reading order.** Shapes are sorted by position, top then left, with a row
  tolerance (3% of the slide height) so side-by-side boxes read left to right,
  rather than by z-order. A group is placed by its own position and its
  children are ordered the same way inside it; child positions are mapped
  through the group transform to slide coordinates.
- **Bullets.** Each paragraph's bullet comes from the first definition found
  in: its own ``a:pPr``, the shape's list style, the layout and master
  placeholders it inherits from, then the master text styles (``bodyStyle``
  for placeholders, ``otherStyle`` for text boxes). So body placeholders are
  bulleted by default, text boxes are not unless they say so, and
  ``a:buNone`` turns bullets off. ``a:buChar`` renders ``-``;
  ``a:buAutoNum`` renders ``1.``, ``a)``, ``iv.``... honouring ``startAt``.
  ``paragraph.level`` gives the nesting (two spaces per level).
- **Tables** become GFM tables (merged cells stay empty). **Charts** become
  ``**Chart: <title>** (<type>)`` plus a table of categories by series.
  **SmartArt** text is listed as bullets. **Speaker notes** follow under
  ``### Notes``.
- **Images** are lazy records with ids ``s<slide>-<n>`` in reading order and
  alt text from ``p:cNvPr/@descr``, with a ``![alt](image:s1-1)`` marker where
  they sit. Images under the size floor are dropped; the marker stays only
  when they have alt text.
"""

from __future__ import annotations

import logging
from collections.abc import Iterator
from dataclasses import dataclass
from typing import Any

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.shapes.picture import Picture
from pptx.text.text import _Run

from ..models.document import RawDocument
from .base import BaseExtractor
from .image_utils import image_marker, image_record
from .markdown_utils import (
    ListDepth,
    MarkdownBuilder,
    Run,
    escape_cell,
    list_item,
    list_marker,
    one_line,
    render_runs,
    render_table,
)

logger = logging.getLogger(__name__)

EMU_PER_POINT = 12700

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
A = f"{{{A_NS}}}"
P = f"{{{P_NS}}}"
DGM_URI = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

_TITLE_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE}
_SKIPPED_PLACEHOLDERS = {PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER}
_BULLET_TAGS = {f"{A}buNone", f"{A}buChar", f"{A}buAutoNum", f"{A}buBlip"}

#: ``a:buAutoNum/@type`` prefix -> list format, suffix -> punctuation.
_AUTONUM_FORMATS = [
    ("arabic", "decimal"),
    ("alphaLc", "lowerLetter"),
    ("alphaUc", "upperLetter"),
    ("romanLc", "lowerRoman"),
    ("romanUc", "upperRoman"),
]

#: Fraction of the slide height within which shapes count as one row.
ROW_TOLERANCE = 0.03


@dataclass
class _Bullet:
    kind: str  # "none" | "bullet" | "number"
    fmt: str = "decimal"
    suffix: str = "."
    start: int = 1


@dataclass
class _Placed:
    """A shape with its slide-space bounding box (EMU)."""

    shape: Any
    left: int
    top: int
    width: int
    height: int
    in_group: bool


def _autonum(scheme: str) -> tuple[str, str]:
    fmt = next((f for prefix, f in _AUTONUM_FORMATS if scheme.startswith(prefix)), "decimal")
    suffix = ")" if "Paren" in scheme else "."
    return fmt, suffix


def _bullet_from(ppr: Any) -> _Bullet | None:
    """The bullet a ``a:pPr``/``a:lvlNpPr`` element defines, if it defines one."""
    if ppr is None:
        return None
    for child in ppr:
        if child.tag not in _BULLET_TAGS:
            continue
        if child.tag == f"{A}buNone":
            return _Bullet("none")
        if child.tag == f"{A}buAutoNum":
            fmt, suffix = _autonum(child.get("type") or "arabicPeriod")
            start = int(child.get("startAt") or 1)
            return _Bullet("number", fmt, suffix, start)
        return _Bullet("bullet")
    return None


def _number(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    return str(value)


class PPTXExtractor(BaseExtractor):
    """Extract PPTX content as Markdown (see the module docstring for the rules)."""

    SUPPORTED_EXTENSIONS = [".pptx", ".ppt"]

    def __init__(
        self,
        min_image_width: int = 100,
        min_image_height: int = 100,
        extract_notes: bool = True,
        image_helper: Any = None,
    ):
        """
        Initialize PPTX extractor.

        Args:
            min_image_width: Minimum image width to keep
            min_image_height: Minimum image height to keep
            extract_notes: Include speaker notes
            image_helper: Unused; kept for API compatibility
        """
        self.min_image_width = min_image_width
        self.min_image_height = min_image_height
        self.extract_notes = extract_notes
        self.image_helper = image_helper

    # ------------------------------------------------------------------ entry

    def extract(self, file_path: str) -> RawDocument:
        """Extract a PPTX file into Markdown text, metadata and per-slide structure."""
        self.validate_file(file_path)
        with open(file_path, "rb") as fh:
            legacy = fh.read(4) == b"\xd0\xcf\x11\xe0"
        if file_path.lower().endswith(".ppt") or legacy:
            raise ValueError(
                "Legacy binary .ppt files can't be read; save the deck as .pptx first."
            )
        prs = Presentation(file_path)
        self._slide_height = int(prs.slide_height or 6858000)

        slides: list[dict[str, Any]] = []
        texts: list[str] = []
        hidden: list[int] = []
        for number, slide in enumerate(prs.slides, start=1):
            data, text = self._slide(slide, number)
            slides.append(data)
            texts.append(text)
            if data.get("hidden"):
                hidden.append(number)

        metadata = self._get_basic_metadata(file_path)
        metadata.update(self._extract_pptx_metadata(prs))
        if hidden:
            metadata["hidden_slides"] = hidden

        return RawDocument(text="\n\n".join(texts), metadata=metadata, structure={"slides": slides})

    # ------------------------------------------------------------------ slide

    def _slide(self, slide: Any, number: int) -> tuple[dict[str, Any], str]:
        self._slide_obj = slide
        self._slide_number = number
        self._image_count = 0
        content: list[dict[str, Any]] = []
        is_hidden = slide._element.get("show") in ("0", "false")

        title_shape = None
        for shape in slide.placeholders:
            if shape.placeholder_format.type in _TITLE_TYPES and shape.has_text_frame:
                title_shape = shape
                break
        title = one_line(title_shape.text_frame.text) if title_shape is not None else ""

        header = f"## Slide {number}" + (f": {title}" if title else "")
        if is_hidden:
            header += " (hidden)"
        builder = MarkdownBuilder()
        builder.add_block(header)
        if title and title_shape is not None:
            content.append({"type": "title", "text": title_shape.text_frame.text})

        top = [s for s in slide.shapes if title_shape is None or s.shape_id != title_shape.shape_id]
        for placed in self._ordered(self._place(top, None)):
            for item, md in self._shape_items(placed):
                content.append(item)
                builder.add_block(md)

        if self.extract_notes and slide.has_notes_slide:
            notes = self._text_frame(slide.notes_slide.notes_text_frame, None, inherit=False)
            if notes.strip():
                content.append({"type": "notes", "text": notes})
                builder.add_block("### Notes")
                builder.add_block(notes)

        data: dict[str, Any] = {
            "slide_number": number,
            "layout": slide.slide_layout.name,
            "content": content,
        }
        if is_hidden:
            data["hidden"] = True
        return data, builder.build()

    # -------------------------------------------------------- reading order

    def _place(self, shapes: Any, transform: tuple[float, float, float, float] | None) -> list[_Placed]:
        """Bounding boxes in slide space; ``transform`` maps group child space."""
        placed = []
        for shape in shapes:
            left, top = shape.left or 0, shape.top or 0
            width, height = shape.width or 0, shape.height or 0
            if transform is not None:
                ox, oy, sx, sy = transform
                left, top = int(ox + left * sx), int(oy + top * sy)
                width, height = int(width * sx), int(height * sy)
            placed.append(_Placed(shape, left, top, width, height, transform is not None))
        return placed

    def _ordered(self, placed: list[_Placed]) -> list[_Placed]:
        """Top-to-bottom rows, left-to-right within a row."""
        tolerance = self._slide_height * ROW_TOLERANCE
        rows: list[list[_Placed]] = []
        for item in sorted(placed, key=lambda p: (p.top, p.left)):
            if rows and item.top - rows[-1][0].top <= tolerance:
                rows[-1].append(item)
            else:
                rows.append([item])
        return [item for row in rows for item in sorted(row, key=lambda p: p.left)]

    def _group_transform(self, group: _Placed) -> tuple[float, float, float, float]:
        """``(offset x, offset y, scale x, scale y)`` from child space to slide space."""
        xfrm = group.shape._element.find(f"{P}grpSpPr/{A}xfrm")
        ch_off = xfrm.find(f"{A}chOff") if xfrm is not None else None
        ch_ext = xfrm.find(f"{A}chExt") if xfrm is not None else None
        cx = int(ch_off.get("x")) if ch_off is not None else 0
        cy = int(ch_off.get("y")) if ch_off is not None else 0
        cw = int(ch_ext.get("cx")) if ch_ext is not None else 0
        ch = int(ch_ext.get("cy")) if ch_ext is not None else 0
        sx = group.width / cw if cw else 1.0
        sy = group.height / ch if ch else 1.0
        return group.left - cx * sx, group.top - cy * sy, sx, sy

    # ----------------------------------------------------------------- shapes

    def _shape_items(self, placed: _Placed) -> Iterator[tuple[dict[str, Any], str]]:
        """``(structure item, markdown)`` pairs for one shape, in order."""
        shape = placed.shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            children = self._place(shape.shapes, self._group_transform(placed))
            for child in self._ordered(children):
                yield from self._shape_items(child)
            return
        # One odd shape (linked picture, OLE object, broken XML) must not take
        # the rest of the slide -- or the deck -- down with it.
        try:
            yield from self._leaf_items(placed)
        except (AttributeError, KeyError, ValueError, TypeError, IndexError) as e:
            logger.warning(
                f"Skipping shape {shape.shape_id} ({shape.name!r}) on slide {self._slide_number}: {e}"
            )

    def _leaf_items(self, placed: _Placed) -> Iterator[tuple[dict[str, Any], str]]:
        shape = placed.shape
        if shape.is_placeholder and shape.placeholder_format.type in _SKIPPED_PLACEHOLDERS:
            return
        # Pictures, including pictures dropped into a layout's picture
        # placeholder (PlaceholderPicture subclasses Picture, but its
        # shape_type is PLACEHOLDER, not PICTURE -- so check the class).
        if isinstance(shape, Picture):
            yield from self._picture(placed)
            return
        if getattr(shape, "has_table", False) and shape.has_table:
            md = self._table(shape.table)
            if md:
                yield {"type": "table", "text": md}, md
            return
        if getattr(shape, "has_chart", False) and shape.has_chart:
            md = self._chart(shape.chart)
            yield {"type": "chart", "text": md}, md
            return
        smartart = self._smartart(shape)
        if smartart:
            yield {"type": "text", "text": smartart}, smartart
            return
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            md = self._text_frame(shape.text_frame, shape)
            if md.strip():
                yield {"type": "text", "text": md}, md

    def _picture(self, placed: _Placed) -> Iterator[tuple[dict[str, Any], str]]:
        shape = placed.shape
        self._image_count += 1
        image_id = f"s{self._slide_number}-{self._image_count}"
        c_nv = shape._element.find(f".//{P}cNvPr")
        alt = ((c_nv.get("descr") or c_nv.get("title")) if c_nv is not None else "") or ""
        alt = alt.strip()
        extra: dict[str, Any] = {"id": image_id, "alt": alt, "page": self._slide_number, "slide": self._slide_number}
        if placed.width and placed.height:
            extra["position"] = {
                "x0": placed.left / EMU_PER_POINT,
                "y0": placed.top / EMU_PER_POINT,
                "x1": (placed.left + placed.width) / EMU_PER_POINT,
                "y1": (placed.top + placed.height) / EMU_PER_POINT,
            }
        try:
            blob = shape.image.blob
        except (AttributeError, KeyError, ValueError) as e:  # linked or missing image
            logger.warning(f"Image {image_id} has no embedded data: {e}")
            blob = None
        record = image_record(blob, self.min_image_width, self.min_image_height, **extra) if blob else None
        marker = image_marker(image_id, alt)
        if record is not None:
            yield record, marker
        elif alt:
            yield {"type": "text", "text": marker}, marker

    # ------------------------------------------------------------------- text

    def _text_frame(self, frame: Any, shape: Any, inherit: bool = True) -> str:
        """Paragraphs as Markdown, with inherited bullets and numbering."""
        builder = MarkdownBuilder()
        depth = ListDepth()
        counters: dict[int, tuple[str, int]] = {}  # level -> (scheme, last number)
        for paragraph in frame.paragraphs:
            text = self._paragraph_text(paragraph)
            if not text.strip():
                continue
            level = paragraph.level
            bullet = self._bullet(paragraph, level, shape, inherit)
            for deeper in [k for k in counters if k > level]:
                del counters[deeper]
            if bullet.kind == "none":
                counters.pop(level, None)
                depth.reset()
                builder.add_block(text.strip())
                continue
            if bullet.kind == "number":
                scheme = f"{bullet.fmt}{bullet.suffix}"
                prev = counters.get(level)
                value = prev[1] + 1 if prev and prev[0] == scheme else bullet.start
                counters[level] = (scheme, value)
                marker = list_marker(bullet.fmt, value, bullet.suffix)
            else:
                counters.pop(level, None)
                marker = "-"
            builder.add_list_item(list_item(text.strip(), depth.depth(level), marker))
        return builder.build()

    def _paragraph_text(self, paragraph: Any) -> str:
        runs: list[Run] = []
        for child in paragraph._p:
            if child.tag == f"{A}r":
                rpr = child.find(f"{A}rPr")
                bold = rpr is not None and rpr.get("b") in ("1", "true")
                italic = rpr is not None and rpr.get("i") in ("1", "true")
                strike = rpr is not None and rpr.get("strike") not in (None, "noStrike")
                href = None
                if rpr is not None and rpr.find(f"{A}hlinkClick") is not None:
                    try:
                        href = _Run(child, paragraph).hyperlink.address
                    except (KeyError, AttributeError):
                        href = None
                runs.append(Run(child.findtext(f"{A}t") or "", bold, italic, href, strike=strike))
            elif child.tag == f"{A}br":
                runs.append(Run("\n"))
            elif child.tag == f"{A}fld":
                runs.append(Run(child.findtext(f"{A}t") or ""))
        return render_runs(runs).replace("\v", "\n")

    def _bullet(self, paragraph: Any, level: int, shape: Any, inherit: bool) -> _Bullet:
        own = _bullet_from(paragraph._p.find(f"{A}pPr"))
        if own is not None or not inherit or shape is None:
            return own or _Bullet("none")
        for lvl_ppr in self._inherited_levels(shape, level):
            found = _bullet_from(lvl_ppr)
            if found is not None:
                return found
        return _Bullet("none")

    def _inherited_levels(self, shape: Any, level: int) -> Iterator[Any]:
        """``a:lvlNpPr`` elements a shape's paragraphs inherit from, nearest first."""
        tag = f"{A}lvl{level + 1}pPr"
        yield self._list_style_level(shape._element, tag)
        slide = self._slide_obj
        master = slide.slide_layout.slide_master
        style = "otherStyle"
        if shape.is_placeholder:
            ph = shape.placeholder_format
            style = "titleStyle" if ph.type in _TITLE_TYPES else "bodyStyle"
            layout_ph = self._matching_placeholder(slide.slide_layout.placeholders, ph)
            if layout_ph is not None:
                yield self._list_style_level(layout_ph._element, tag)
            master_ph = self._matching_placeholder(master.placeholders, ph, by_type_only=True)
            if master_ph is not None:
                yield self._list_style_level(master_ph._element, tag)
        styles = master._element.find(f"{P}txStyles/{P}{style}")
        if styles is not None:
            yield styles.find(tag)

    @staticmethod
    def _list_style_level(element: Any, tag: str) -> Any:
        return element.find(f".//{A}lstStyle/{tag}")

    @staticmethod
    def _matching_placeholder(placeholders: Any, ph: Any, by_type_only: bool = False) -> Any:
        body_like = {PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT}
        for candidate in placeholders:
            fmt = candidate.placeholder_format
            if not by_type_only and fmt.idx == ph.idx and ph.idx != 0:
                return candidate
        for candidate in placeholders:
            fmt = candidate.placeholder_format
            if fmt.type == ph.type or (fmt.type in body_like and ph.type in body_like):
                return candidate
        return None

    # ------------------------------------------------------- tables & charts

    def _table(self, table: Any) -> str:
        rows: list[list[str]] = []
        for row in table.rows:
            cells = []
            for cell in row.cells:
                if cell.is_spanned:
                    cells.append("")
                    continue
                lines = [self._paragraph_text(p).strip() for p in cell.text_frame.paragraphs]
                cells.append(escape_cell("\n".join(line for line in lines if line)))
            rows.append(cells)
        return render_table(rows)

    def _chart(self, chart: Any) -> str:
        title = ""
        if chart.has_title and chart.chart_title.has_text_frame:
            title = one_line(chart.chart_title.text_frame.text)
        try:
            kind = chart.chart_type.name.replace("_", " ").lower()
        except (AttributeError, ValueError, KeyError, NotImplementedError):
            kind = ""
        head = f"**Chart: {title}**" if title else "**Chart**"
        if kind:
            head += f" ({kind})"
        try:
            series = list(chart.series)
            categories = [str(c) for c in chart.plots[0].categories] if len(chart.plots) else []
            values = [list(s.values) for s in series]
        except (AttributeError, ValueError, KeyError, IndexError, TypeError) as e:
            logger.warning(f"Chart data on slide {self._slide_number} could not be read: {e}")
            return head
        if not series:
            return head
        count = max((len(v) for v in values), default=0)
        labels = categories if len(categories) == count else [str(i + 1) for i in range(count)]
        rows = [["Category", *[s.name or f"Series {i + 1}" for i, s in enumerate(series)]]]
        for i, label in enumerate(labels):
            rows.append([label, *[_number(v[i]) if i < len(v) else "" for v in values]])
        return f"{head}\n\n{render_table(rows)}"

    def _smartart(self, shape: Any) -> str:
        """Text of a SmartArt diagram (its data model part), one bullet per node."""
        graphic_data = shape._element.find(f".//{A}graphicData")
        if graphic_data is None or graphic_data.get("uri") != DGM_URI:
            return ""
        rel_ids = graphic_data.find(f"{{{DGM_URI}}}relIds")
        rid = rel_ids.get(f"{{{R_NS}}}dm") if rel_ids is not None else None
        if not rid:
            return ""
        try:
            part = self._slide_obj.part.related_part(rid)
            root = etree.fromstring(part.blob, parser=etree.XMLParser(resolve_entities=False, no_network=True))
        except (KeyError, etree.XMLSyntaxError) as e:
            logger.warning(f"SmartArt on slide {self._slide_number} could not be read: {e}")
            return ""
        lines = []
        for pt in root.iter(f"{{{DGM_URI}}}pt"):
            if pt.get("type") not in (None, "node"):
                continue
            text = one_line(" ".join(t.text or "" for t in pt.iter(f"{A}t")))
            if text:
                lines.append(f"- {text}")
        return "\n".join(lines)

    # --------------------------------------------------------------- metadata

    def _extract_pptx_metadata(self, prs: Any) -> dict[str, Any]:
        core_props = prs.core_properties
        return {
            "author": core_props.author or "Unknown",
            "title": core_props.title or "",
            "subject": core_props.subject or "",
            "created": str(core_props.created) if core_props.created else "",
            "modified": str(core_props.modified) if core_props.modified else "",
            "slide_count": len(prs.slides),
            "format": "pptx",
        }
