"""
Exact-output tests for the Excel and CSV extractors.

Every fixture is a real file written into ``tmp_path`` (openpyxl workbooks,
CSV text) and every assertion compares the full Markdown the extractor emits,
so a regression in spacing, header choice or number formatting shows up.

Run: ``python -m pytest src/tests/test_fidelity_tabular.py -q``
"""

from __future__ import annotations

import sys
import warnings
from datetime import date, datetime
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

openpyxl = pytest.importorskip("openpyxl")
pytest.importorskip("pandas")

from openpyxl.styles import Font, PatternFill  # noqa: E402

from src.chunking.strategies import ExcelChunker  # noqa: E402
from src.extractors.csv_extractor import CSVExtractor  # noqa: E402
from src.extractors.excel_extractor import ExcelExtractor  # noqa: E402

# ------------------------------------------------------------------------ Excel


def _save(wb: openpyxl.Workbook, tmp_path: Path, name: str = "book.xlsx") -> str:
    path = tmp_path / name
    wb.save(path)
    wb.close()
    return str(path)


def test_excel_merged_title_dates_and_number_formats(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Report"
    ws["A1"] = "Quarterly Sales"
    ws.merge_cells("A1:D1")
    ws.append(["Region", "Date", "Growth", "Revenue"])
    ws.append(["North", date(2024, 3, 1), 0.125, 1234567])
    ws.append(["South", datetime(2024, 3, 2, 14, 30), 0.05, 987.5])
    ws["C3"].number_format = "0.0%"
    ws["C4"].number_format = "0%"
    ws["D3"].number_format = "#,##0"
    ws["D4"].number_format = "#,##0.00"
    raw = ExcelExtractor().extract(_save(wb, tmp_path))

    assert raw.text == (
        "## Sheet: Report\n"
        "\n"
        "### Quarterly Sales\n"
        "\n"
        "| Region | Date | Growth | Revenue |\n"
        "| --- | --- | --- | --- |\n"
        "| North | 2024-03-01 | 12.5% | 1,234,567 |\n"
        "| South | 2024-03-02 14:30 | 5% | 987.50 |"
    )
    sheet = raw.structure["sheets"][0]
    assert sheet["markdown"] == raw.text
    table = sheet["info"].tables[0]
    assert (table.start_row, table.header_row, table.end_row) == (1, 2, 4)


def test_excel_formatting_only_cells_do_not_blow_up(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Name", "Score"])
    ws.append(["Ann", 0.30000000000000004])
    fill = PatternFill("solid", fgColor="FFFF00")
    for row in range(1, 5001, 7):
        for col in range(1, 51, 3):
            ws.cell(row, col).fill = fill
    ws.cell(5000, 50).fill = fill
    raw = ExcelExtractor().extract(_save(wb, tmp_path))

    expected = "| Name | Score |\n| --- | --- |\n| Ann | 0.3 |"
    assert raw.text == "## Sheet: Data\n\n" + expected
    # The row-aligned grid kept for the chunker is bounded the same way.
    assert raw.structure["sheets"][0]["text"] == expected
    info = raw.structure["sheets"][0]["info"]
    assert (info.max_row, info.max_column) == (2, 2)


def test_excel_hidden_sheet_and_uncached_formula(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    for row in (["Item", "Qty"], ["x", 1], ["y", 2], ["Total", "=SUM(B2:B3)"]):
        ws.append(row)
    secret = wb.create_sheet("Secret")
    secret.append(["Key", "Value"])
    secret.append(["pin", 1234])
    secret.sheet_state = "hidden"
    raw = ExcelExtractor().extract(_save(wb, tmp_path))

    assert raw.text == (
        "## Sheet: Data\n"
        "\n"
        "| Item | Qty |\n"
        "| --- | --- |\n"
        "| x | 1 |\n"
        "| y | 2 |\n"
        "| Total | =SUM(B2:B3) |\n"
        "\n"
        "## Sheet: Secret (hidden)\n"
        "\n"
        "| Key | Value |\n"
        "| --- | --- |\n"
        "| pin | 1234 |"
    )
    assert raw.metadata["hidden_sheets"] == ["Secret"]
    assert raw.metadata["num_sheets"] == 2
    assert raw.structure["excel_metadata"].sheets[0].has_formulas


def test_excel_separate_tables_and_vertical_merge(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Teams"
    ws.append(["Name", "Team"])
    ws.append(["Ann", "Red"])
    ws.append(["Bob", None])
    ws.merge_cells("B2:B3")
    ws.append([])
    ws.append(["Code", "Count"])
    ws.append(["A", 3])
    raw = ExcelExtractor().extract(_save(wb, tmp_path))

    assert raw.text == (
        "## Sheet: Teams\n"
        "\n"
        "| Name | Team |\n"
        "| --- | --- |\n"
        "| Ann | Red |\n"
        "| Bob |  |\n"
        "\n"
        "| Code | Count |\n"
        "| --- | --- |\n"
        "| A | 3 |"
    )
    grid = raw.structure["sheets"][0]["text"]
    assert grid == (
        "| Name | Team |\n"
        "| --- | --- |\n"
        "| Ann | Red |\n"
        "| Bob |  |\n"
        "|  |  |\n"
        "| Code | Count |\n"
        "| A | 3 |"
    )
    # The grid stays row-aligned, so the chunker's by-table slicing is right.
    chunks = ExcelChunker(strategy="by_table", convert_to_html=False).chunk(raw)
    assert [c.text for c in chunks] == ["Name\tTeam\nAnn\tRed\nBob\t", "Code\tCount\nA\t3"]


def test_excel_numeric_first_row_is_not_a_header(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Numbers"
    ws.append([1, 2])
    ws.append([3, 4])
    raw = ExcelExtractor().extract(_save(wb, tmp_path))
    assert raw.text == "## Sheet: Numbers\n\n|  |  |\n| --- | --- |\n| 1 | 2 |\n| 3 | 4 |"


def test_excel_bold_header_after_unmerged_title(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws["A1"] = "Budget 2025"
    ws.append(["Line", "Amount"])
    ws.append(["Rent", 1200])
    for cell in ws[2]:
        cell.font = Font(bold=True)
    raw = ExcelExtractor().extract(_save(wb, tmp_path))
    assert raw.text == (
        "## Sheet: Budget\n\n### Budget 2025\n\n| Line | Amount |\n| --- | --- |\n| Rent | 1200 |"
    )


# -------------------------------------------------------------------------- CSV


def _csv(tmp_path: Path, content: str | bytes, name: str = "data.csv") -> str:
    path = tmp_path / name
    if isinstance(content, bytes):
        path.write_bytes(content)
    else:
        path.write_text(content, encoding="utf-8")
    return str(path)


def test_csv_bad_lines_are_kept_and_reported(tmp_path: Path) -> None:
    path = _csv(
        tmp_path,
        'name,qty,note\napple,3,fresh\npear,4\nplum,5,ripe,extra\nfig,6,"two\nlines"\n',
    )
    raw = CSVExtractor().extract(path)
    assert raw.text == (
        "| name | qty | note |\n"
        "| --- | --- | --- |\n"
        "| apple | 3 | fresh |\n"
        "| pear | 4 |  |\n"
        "| plum | 5 | ripe,extra |\n"
        "| fig | 6 | two<br>lines |"
    )
    assert raw.metadata["bad_lines"] == [3, 4]
    assert raw.metadata["bad_line_count"] == 2
    assert raw.metadata["num_rows"] == 4


def test_csv_ints_with_missing_values_stay_ints(tmp_path: Path) -> None:
    raw = CSVExtractor().extract(_csv(tmp_path, "id,score\n1,10\n2,\n3,30\n"))
    assert raw.text == "| id | score |\n| --- | --- |\n| 1 | 10 |\n| 2 |  |\n| 3 | 30 |"
    assert raw.metadata["column_types"] == {"id": "integer", "score": "integer"}
    assert "bad_lines" not in raw.metadata or raw.metadata["bad_lines"] == []


def test_csv_semicolon_with_preamble(tmp_path: Path) -> None:
    content = (
        "Monthly export\n"
        "generated 2024-01-05\n"
        "\n"
        "date;amount;city\n"
        "2024-01-01;1.234,5;Paris\n"
        "2024-01-02;99;Lyon\n"
    )
    raw = CSVExtractor().extract(_csv(tmp_path, content))
    assert raw.metadata["delimiter"] == ";"
    assert raw.metadata["preamble"] == ["Monthly export", "generated 2024-01-05"]
    assert raw.text == (
        "| date | amount | city |\n"
        "| --- | --- | --- |\n"
        "| 2024-01-01 | 1.234,5 | Paris |\n"
        "| 2024-01-02 | 99 | Lyon |"
    )
    assert raw.metadata["column_names"] == ["date", "amount", "city"]


def test_csv_excel_sep_hint_line(tmp_path: Path) -> None:
    raw = CSVExtractor().extract(_csv(tmp_path, "sep=|\na|b\n1|2\n"))
    assert raw.metadata["delimiter"] == "|"
    assert raw.text == "| a | b |\n| --- | --- |\n| 1 | 2 |"


def test_csv_dates_infer_without_warnings(tmp_path: Path) -> None:
    content = "when,what\n2024-01-01,a\n01/02/2024,b\nMarch 3 2024,c\n"
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        raw = CSVExtractor().extract(_csv(tmp_path, content))
    assert [w for w in caught if issubclass(w.category, UserWarning)] == []
    assert raw.metadata["column_types"]["when"] == "datetime"
    assert raw.text.splitlines()[3] == "| 01/02/2024 | b |"


def test_csv_without_header_and_legacy_encoding(tmp_path: Path) -> None:
    raw = CSVExtractor().extract(_csv(tmp_path, "1,2\n3,4\n"))
    assert raw.text == "|  |  |\n| --- | --- |\n| 1 | 2 |\n| 3 | 4 |"
    assert raw.metadata["has_header"] is False

    latin = CSVExtractor().extract(
        _csv(tmp_path, "city,note\nZürich,café crème\n".encode("cp1252"), "latin.csv")
    )
    assert latin.text == "| city | note |\n| --- | --- |\n| Zürich | café crème |"


def test_csv_tab_delimited(tmp_path: Path) -> None:
    raw = CSVExtractor().extract(_csv(tmp_path, "Name\tAge\nJohn\t30\n", "people.tsv"))
    assert raw.metadata["delimiter"] == "\t"
    assert raw.text == "| Name | Age |\n| --- | --- |\n| John | 30 |"
