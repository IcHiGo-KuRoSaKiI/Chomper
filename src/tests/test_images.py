"""
Image extraction, end to end, for PDF, PPTX and DOCX.

Every entry point that hands images back -- ``chomper.parse``,
``chomper.extract_metadata``, and the MCP ``get_document_images`` /
``parse_document`` / ``extract_metadata`` tools -- must see the same images.
Each extractor nests images differently, and callers reading only one shape is
how images were silently dropped (issue #1 for PDF, then PPTX). Fixtures are
built in ``tmp_path`` so the tests need nothing checked in.

Run: ``python -m pytest src/tests/test_images.py -q``
"""

from __future__ import annotations

import asyncio
import base64
import io
import json
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

fitz = pytest.importorskip("fitz", reason="PyMuPDF required")
pptx = pytest.importorskip("pptx", reason="python-pptx required")
docx = pytest.importorskip("docx", reason="python-docx required")
from PIL import Image, ImageDraw  # noqa: E402
from pptx.util import Inches  # noqa: E402

import chomper  # noqa: E402
from src.handlers.images import handle_get_document_images  # noqa: E402
from src.handlers.metadata import handle_extract_metadata  # noqa: E402
from src.handlers.parse import handle_parse_document  # noqa: E402

CLIENT_MIMES = {"image/png", "image/jpeg", "image/gif", "image/webp"}


def _png(w: int, h: int, color: str, transparent: bool = False) -> bytes:
    img = Image.new("RGBA" if transparent else "RGB", (w, h), (0, 0, 0, 0) if transparent else "white")
    ImageDraw.Draw(img).ellipse([10, 10, w - 10, h - 10], fill=color)
    out = io.BytesIO()
    img.save(out, "PNG")
    return out.getvalue()


def _decode(b64: str) -> Image.Image:
    return Image.open(io.BytesIO(base64.b64decode(b64)))


def _mcp_images(path: Path, **args) -> list:
    reply = asyncio.run(handle_get_document_images({"file_path": str(path), "max_images": 50, **args}))
    return [item for item in reply if item.type == "image"]


# --------------------------------------------------------------------------- PDF


@pytest.fixture
def pdf_path(tmp_path: Path) -> Path:
    """Three pages: a logo on every page, a transparent PNG, a CMYK JPEG."""
    doc = fitz.open()
    logo = _png(120, 120, "purple")
    cmyk = io.BytesIO()
    Image.new("CMYK", (300, 200), (0, 255, 255, 0)).save(cmyk, "JPEG")
    for i in range(3):
        page = doc.new_page()
        page.insert_text((72, 72), f"Page {i + 1}")
        page.insert_image(fitz.Rect(450, 20, 560, 130), stream=logo)
        if i == 0:
            page.insert_image(fitz.Rect(72, 150, 372, 375), stream=_png(400, 300, "red", transparent=True))
        if i == 2:
            page.insert_image(fitz.Rect(72, 150, 372, 350), stream=cmyk.getvalue())
    path = tmp_path / "images.pdf"
    doc.save(path)
    return path


def test_pdf_images_reach_every_entry_point(pdf_path: Path):
    parsed = chomper.parse(str(pdf_path), include_images=True).images
    assert len(parsed) == 3  # logo once + transparent PNG + CMYK JPEG
    assert chomper.extract_metadata(str(pdf_path)).image_count == 3
    assert len(_mcp_images(pdf_path)) == 3


def test_pdf_repeated_logo_is_deduped_but_page_filter_keeps_it(pdf_path: Path):
    parsed = chomper.parse(str(pdf_path), include_images=True).images
    logo = next(img for img in parsed if img["width"] == 120)
    assert logo["pages"] == [1, 2, 3]
    # Asking for one page still shows everything on that page, logo included.
    assert len(_mcp_images(pdf_path, page=2)) == 1
    assert len(_mcp_images(pdf_path, page=1)) == 2


def test_pdf_transparency_survives(pdf_path: Path):
    parsed = chomper.parse(str(pdf_path), include_images=True).images
    red = _decode(next(img for img in parsed if img["width"] == 400)["content"])
    assert "A" in red.getbands(), "soft mask was dropped"
    assert red.getpixel((0, 0))[3] == 0, "transparent corner came out opaque"


def test_pdf_images_are_client_safe(pdf_path: Path):
    for img in chomper.parse(str(pdf_path), include_images=True).images:
        assert img["mime_type"] in CLIENT_MIMES
        decoded = _decode(img["content"])
        assert decoded.mode != "CMYK"
        assert Image.MIME[decoded.format] == img["mime_type"]


# -------------------------------------------------------------------------- PPTX


@pytest.fixture
def pptx_path(tmp_path: Path) -> Path:
    """A realistic deck: titles, text boxes, a table, grouped and placeholder pictures."""
    prs = pptx.Presentation()

    s1 = prs.slides.add_slide(prs.slide_layouts[5])  # title only
    s1.shapes.title.text = "Quarterly numbers"
    s1.shapes.add_textbox(Inches(1), Inches(2), Inches(3), Inches(1)).text_frame.text = "a plain text box"
    s1.shapes.add_picture(io.BytesIO(_png(400, 300, "red")), Inches(4), Inches(2))
    s1.shapes.add_table(2, 2, Inches(1), Inches(4), Inches(3), Inches(1)).table.cell(0, 0).text = "cell"

    s2 = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    group = s2.shapes.add_group_shape()
    group.shapes.add_picture(io.BytesIO(_png(400, 300, "blue")), Inches(1), Inches(1))
    group.shapes.add_textbox(Inches(1), Inches(5), Inches(3), Inches(1)).text_frame.text = "grouped caption"

    s3 = prs.slides.add_slide(prs.slide_layouts[8])  # picture with caption
    s3.shapes.title.text = "Placeholder slide"
    holder = next(p for p in s3.placeholders if p.placeholder_format.type == 18)
    holder.insert_picture(io.BytesIO(_png(400, 300, "green")))

    s4 = prs.slides.add_slide(prs.slide_layouts[6])
    s4.shapes.add_picture(io.BytesIO(_png(400, 300, "red")), Inches(1), Inches(1))  # same as slide 1
    s4.shapes.add_picture(io.BytesIO(_png(40, 40, "orange")), Inches(6), Inches(1))  # icon, below floor

    path = tmp_path / "deck.pptx"
    prs.save(path)
    return path


def test_pptx_with_text_boxes_does_not_crash(pptx_path: Path):
    result = chomper.parse(str(pptx_path))
    assert "Quarterly numbers" in result.text
    assert "a plain text box" in result.text
    assert "grouped caption" in result.text
    assert "cell" in result.text


def test_pptx_titles_detected_safely(pptx_path: Path):
    from src.extractors.pptx_extractor import PPTXExtractor

    slides = PPTXExtractor().extract(str(pptx_path)).structure["slides"]
    first = slides[0]["content"]
    assert {"type": "title", "text": "Quarterly numbers"} in first
    assert {"type": "text", "text": "a plain text box"} in first


def test_pptx_images_reach_every_entry_point(pptx_path: Path):
    parsed = chomper.parse(str(pptx_path), include_images=True).images
    # red (slides 1 + 4, deduped), grouped blue, placeholder green; icon skipped
    assert len(parsed) == 3
    assert {img["page"] for img in parsed} == {1, 2, 3}
    red = next(img for img in parsed if img["page"] == 1)
    assert red["pages"] == [1, 4]
    assert chomper.extract_metadata(str(pptx_path)).image_count == 3
    assert len(_mcp_images(pptx_path)) == 3
    assert len(_mcp_images(pptx_path, page=4)) == 1


def test_pptx_positions_in_points(pptx_path: Path):
    parsed = chomper.parse(str(pptx_path), include_images=True).images
    red = next(img for img in parsed if img["page"] == 1)
    assert red["position"]["x0"] == pytest.approx(4 * 72)
    assert red["position"]["y0"] == pytest.approx(2 * 72)


def test_mcp_parse_and_metadata_count_pptx_images(pptx_path: Path):
    reply = asyncio.run(handle_parse_document({"file_path": str(pptx_path), "include_images": True}))
    assert sum(1 for item in reply if item.type == "image") == 3
    meta = asyncio.run(handle_extract_metadata({"file_path": str(pptx_path)}))
    assert json.loads(meta[0].text)["document_info"]["image_count"] == 3


def test_legacy_ppt_gives_a_clear_error(tmp_path: Path):
    path = tmp_path / "old.ppt"
    path.write_bytes(b"\xd0\xcf\x11\xe0 not a zip")
    with pytest.raises(ValueError, match="pptx"):
        chomper.parse(str(path))


# -------------------------------------------------------------------------- DOCX


def test_docx_images_reach_every_entry_point(tmp_path: Path):
    document = docx.Document()
    document.add_paragraph("Before the picture")
    document.add_picture(io.BytesIO(_png(400, 300, "red")))
    document.add_paragraph("After the picture")
    path = tmp_path / "doc.docx"
    document.save(path)

    parsed = chomper.parse(str(path), include_images=True).images
    assert len(parsed) == 1 and parsed[0]["mime_type"] == "image/png"
    assert chomper.extract_metadata(str(path)).image_count == 1
    assert len(_mcp_images(path)) == 1


# ------------------------------------------------------------------------ shared


def test_normalize_converts_unsupported_formats():
    from src.extractors.image_utils import normalize_image

    tiff = io.BytesIO()
    Image.new("RGB", (50, 40), "red").save(tiff, "TIFF")
    data, width, height, mime = normalize_image(tiff.getvalue())
    assert (width, height, mime) == (50, 40, "image/png")
    assert _decode(base64.b64encode(data).decode()).format == "PNG"
    assert normalize_image(b"definitely not an image") is None
