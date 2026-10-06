"""
Output fidelity of the DOCX, PPTX, HTML, EPUB and email extractors.

Every test builds a real file in ``tmp_path`` (python-docx, python-pptx,
ebooklib, the stdlib email package) and asserts on the exact Markdown that comes
out, because the text is what an LLM reads: a lost bullet level, a flattened
numbered list or a glued word is an accuracy bug even when "the text is there".

Run: ``python -m pytest src/tests/test_extractor_fidelity.py -q``
"""

from __future__ import annotations

import io
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from PIL import Image  # noqa: E402

docx = pytest.importorskip("docx", reason="python-docx required")
from docx.opc.constants import CONTENT_TYPE as CT  # noqa: E402
from docx.opc.constants import RELATIONSHIP_TYPE as RT  # noqa: E402
from docx.opc.packuri import PackURI  # noqa: E402
from docx.opc.part import Part  # noqa: E402
from docx.oxml import parse_xml  # noqa: E402
from docx.oxml.ns import nsdecls  # noqa: E402

from src.extractors.docx_extractor import DOCXExtractor  # noqa: E402
from src.extractors.image_utils import collect_images  # noqa: E402

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _png(w: int, h: int, color: str = "red") -> bytes:
    out = io.BytesIO()
    Image.new("RGB", (w, h), color).save(out, "PNG")
    return out.getvalue()


# --------------------------------------------------------------------- DOCX kit


def _add_abstract(document, abstract_id: int, levels: list[tuple[str, str, int]]) -> None:
    """Append a ``w:abstractNum`` with ``(numFmt, lvlText, start)`` per level."""
    lvls = "".join(
        f'<w:lvl w:ilvl="{i}"><w:start w:val="{start}"/><w:numFmt w:val="{fmt}"/>'
        f'<w:lvlText w:val="{text}"/><w:lvlJc w:val="left"/>'
        f'<w:pPr><w:ind w:left="{720 * (i + 1)}" w:hanging="360"/></w:pPr></w:lvl>'
        for i, (fmt, text, start) in enumerate(levels)
    )
    numbering = document.part.numbering_part.element
    abstract = parse_xml(
        f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{abstract_id}">{lvls}</w:abstractNum>'
    )
    # abstractNum elements must precede num elements.
    first_num = numbering.find(f"{{{W_NS}}}num")
    first_num.addprevious(abstract)


def _add_num(document, num_id: int, abstract_id: int, start_override: dict[int, int] | None = None) -> None:
    overrides = "".join(
        f'<w:lvlOverride w:ilvl="{lvl}"><w:startOverride w:val="{val}"/></w:lvlOverride>'
        for lvl, val in (start_override or {}).items()
    )
    document.part.numbering_part.element.append(
        parse_xml(
            f'<w:num {nsdecls("w")} w:numId="{num_id}"><w:abstractNumId w:val="{abstract_id}"/>'
            f"{overrides}</w:num>"
        )
    )


def _numbered(document, text: str, num_id: int, ilvl: int):
    paragraph = document.add_paragraph(text)
    paragraph._p.get_or_add_pPr().append(
        parse_xml(
            f'<w:numPr {nsdecls("w")}><w:ilvl w:val="{ilvl}"/><w:numId w:val="{num_id}"/></w:numPr>'
        )
    )
    return paragraph


def _docx_text(path: Path) -> str:
    return DOCXExtractor().extract(str(path)).text


@pytest.fixture
def docx_lists(tmp_path: Path) -> Path:
    document = docx.Document()
    document.add_paragraph("Project Plan", style="Title")
    document.add_heading("Goals", level=1)

    _add_abstract(
        document,
        90,
        [("decimal", "%1.", 1), ("lowerLetter", "%2)", 1), ("lowerRoman", "%3.", 1)],
    )
    _add_num(document, 90, 90)
    _numbered(document, "Plan", 90, 0)
    _numbered(document, "Research", 90, 1)
    _numbered(document, "Budget", 90, 1)
    _numbered(document, "Hardware", 90, 2)
    _numbered(document, "Software", 90, 2)
    _numbered(document, "Build", 90, 0)
    _numbered(document, "Prototype", 90, 1)

    document.add_paragraph("Between lists.")
    document.add_paragraph("Apples", style="List Bullet")
    document.add_paragraph("Green", style="List Bullet 2")
    document.add_paragraph("Red", style="List Bullet 2")
    document.add_paragraph("Pears", style="List Bullet")
    document.add_paragraph("First", style="List Number")
    document.add_paragraph("Second", style="List Number")

    document.add_heading("Restarted", level=2)
    _add_num(document, 91, 90, start_override={0: 1})
    _numbered(document, "Again from one", 91, 0)
    _numbered(document, "And two", 91, 0)

    _add_abstract(document, 92, [("upperRoman", "%1.", 4), ("bullet", "", 1)])
    _add_num(document, 92, 92)
    _numbered(document, "Fourth", 92, 0)
    _numbered(document, "a bullet under it", 92, 1)
    _numbered(document, "Fifth", 92, 0)
    path = tmp_path / "lists.docx"
    document.save(path)
    return path


def test_docx_lists_keep_numbering_nesting_and_marker_types(docx_lists: Path):
    assert _docx_text(docx_lists) == (
        "# Project Plan\n\n"
        "## Goals\n\n"
        "1. Plan\n"
        "  a) Research\n"
        "  b) Budget\n"
        "    i. Hardware\n"
        "    ii. Software\n"
        "2. Build\n"
        "  a) Prototype\n\n"
        "Between lists.\n\n"
        "- Apples\n"
        "  - Green\n"
        "  - Red\n"
        "- Pears\n\n"
        "1. First\n"
        "2. Second\n\n"
        "### Restarted\n\n"
        "1. Again from one\n"
        "2. And two\n\n"
        "IV. Fourth\n"
        "  - a bullet under it\n"
        "V. Fifth"
    )


def test_docx_without_title_maps_heading_levels_directly(tmp_path: Path):
    document = docx.Document()
    document.add_paragraph("Preamble before any heading.")
    document.add_heading("Top", level=1)
    document.add_heading("Sub", level=2)
    document.add_paragraph("Body.")
    path = tmp_path / "plain.docx"
    document.save(path)

    raw = DOCXExtractor().extract(str(path))
    assert raw.text == "Preamble before any heading.\n\n# Top\n\n## Sub\n\nBody."
    assert "Introduction" not in raw.text
    sections = raw.structure["sections"]
    assert sections[0]["heading"] == "" and sections[0]["level"] == 0
    assert sections[1]["heading"] == "Top"


@pytest.fixture
def docx_rich(tmp_path: Path) -> Path:
    document = docx.Document()
    body = document.element.body

    # A body-level content control, like a cover page.
    body.insert(
        0,
        parse_xml(
            f"<w:sdt {nsdecls('w')}><w:sdtPr/><w:sdtContent>"
            "<w:p><w:r><w:t>Cover page text</w:t></w:r></w:p>"
            "</w:sdtContent></w:sdt>"
        )
    )

    # Inline formatting and a hyperlink.
    paragraph = document.add_paragraph("Read ")
    paragraph.add_run("this").bold = True
    paragraph.add_run(" ")
    paragraph.add_run("carefully").italic = True
    paragraph.add_run(" at ")
    rid = document.part.relate_to("https://example.com/docs", RT.HYPERLINK, is_external=True)
    paragraph._p.append(
        parse_xml(
            f'<w:hyperlink {nsdecls("w", "r")} r:id="{rid}"><w:r><w:t>the docs</w:t></w:r></w:hyperlink>'
        )
    )
    paragraph.add_run(".")

    # A text box anchored to a paragraph, in both the DrawingML choice and the
    # VML fallback, as Word writes it.
    anchor = document.add_paragraph("Anchor paragraph")
    anchor._p.append(
        parse_xml(
            f"<w:r {nsdecls('w', 'wp', 'a')} "
            'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
            'xmlns:v="urn:schemas-microsoft-com:vml" '
            'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
            '<mc:AlternateContent><mc:Choice Requires="wps"><w:drawing><wp:anchor>'
            '<wp:docPr id="50" name="Text Box 1"/><a:graphic>'
            '<a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
            "<wps:wsp><wps:txbx><w:txbxContent><w:p><w:r><w:t>Boxed note</w:t></w:r></w:p>"
            "</w:txbxContent></wps:txbx></wps:wsp></a:graphicData></a:graphic>"
            "</wp:anchor></w:drawing></mc:Choice><mc:Fallback><w:pict><v:shape><v:textbox>"
            "<w:txbxContent><w:p><w:r><w:t>Boxed note</w:t></w:r></w:p></w:txbxContent>"
            "</v:textbox></v:shape></w:pict></mc:Fallback></mc:AlternateContent></w:r>"
        )
    )

    # A table: header row with a horizontal merge, a vertical merge, a pipe and a
    # line break in a cell, an image in a cell and a nested table.
    table = document.add_table(rows=4, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Region"
    table.cell(0, 2).text = "Notes"
    table.cell(1, 0).merge(table.cell(2, 0)).text = "North"
    table.cell(1, 1).text = "a|b"
    table.cell(1, 2).text = "line one"
    table.cell(1, 2).paragraphs[0].add_run().add_break()
    table.cell(1, 2).paragraphs[0].add_run("line two")
    table.cell(2, 1).text = "Q2"
    pic = table.cell(2, 2).paragraphs[0].add_run().add_picture(io.BytesIO(_png(200, 150, "blue")))
    pic._inline.docPr.set("descr", "Sales chart")
    table.cell(3, 0).text = "South"
    inner = table.cell(3, 1).add_table(rows=2, cols=2)
    inner.cell(0, 0).text = "k"
    inner.cell(0, 1).text = "v"
    inner.cell(1, 0).text = "x"
    inner.cell(1, 1).text = "1"
    table.cell(3, 2).text = "end"

    # Images: one kept with alt text, a tiny one with alt (marker only), a tiny
    # decorative one without alt (dropped entirely).
    document.add_picture(io.BytesIO(_png(300, 200)))
    document.inline_shapes[-1]._inline.docPr.set("descr", "Team photo")
    tiny = document.add_paragraph().add_run().add_picture(io.BytesIO(_png(20, 20, "green")))
    tiny._inline.docPr.set("descr", "Company logo")
    document.add_paragraph().add_run().add_picture(io.BytesIO(_png(10, 10, "black")))

    # Footnote reference plus a footnotes part.
    note_paragraph = document.add_paragraph("A claim")
    note_paragraph._p.append(
        parse_xml(
            f'<w:r {nsdecls("w")}><w:rPr><w:rStyle w:val="FootnoteReference"/></w:rPr>'
            '<w:footnoteReference w:id="1"/></w:r>'
        )
    )
    note_paragraph.add_run(" stands.")
    footnotes = (
        f'<w:footnotes {nsdecls("w")}>'
        '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
        '<w:footnote w:id="1"><w:p><w:r><w:footnoteRef/></w:r>'
        "<w:r><w:t xml:space=\"preserve\"> Source: annual report.</w:t></w:r></w:p></w:footnote>"
        "</w:footnotes>"
    ).encode()
    part = Part(PackURI("/word/footnotes.xml"), CT.WML_FOOTNOTES, footnotes, document.part.package)
    document.part.relate_to(part, RT.FOOTNOTES)

    section = document.sections[0]
    section.header.paragraphs[0].text = "ACME Confidential"
    section.footer.paragraphs[0].text = "Internal use only"

    path = tmp_path / "rich.docx"
    document.save(path)
    return path


def test_docx_rich_content_is_complete_and_in_order(docx_rich: Path):
    raw = DOCXExtractor().extract(str(docx_rich))
    assert raw.text == (
        "Cover page text\n\n"
        "Read **this** *carefully* at [the docs](https://example.com/docs).\n\n"
        "Anchor paragraph\n\n"
        "Boxed note\n\n"
        "| Region |  | Notes |\n"
        "| --- | --- | --- |\n"
        "| North | a\\|b | line one<br>line two |\n"
        "|  | Q2 | ![Sales chart](image:img-1) |\n"
        "| South | k; v<br>x; 1 | end |\n\n"
        "![Team photo](image:img-2)\n\n"
        "![Company logo](image:img-3)\n\n"
        "A claim[^1] stands.\n\n"
        "[^1]: Source: annual report."
    )
    assert raw.metadata["headers"] == ["ACME Confidential"]
    assert raw.metadata["footers"] == ["Internal use only"]
    assert "ACME" not in raw.text


def test_docx_image_records_match_markers(docx_rich: Path):
    raw = DOCXExtractor().extract(str(docx_rich))
    images = collect_images(raw.structure)
    assert [(img["id"], img["alt"], img["width"]) for img in images] == [
        ("img-1", "Sales chart", 200),
        ("img-2", "Team photo", 300),
    ]


def test_legacy_doc_gives_a_clear_error(tmp_path: Path):
    path = tmp_path / "old.doc"
    path.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 512)
    with pytest.raises(ValueError, match=r"legacy binary Word 97-2003 \(\.doc\).*save it as \.docx"):
        DOCXExtractor().extract(str(path))


# ------------------------------------------------------------------------- PPTX

pptx = pytest.importorskip("pptx", reason="python-pptx required")
from pptx.chart.data import CategoryChartData  # noqa: E402
from pptx.enum.chart import XL_CHART_TYPE  # noqa: E402
from pptx.oxml import parse_xml as pptx_parse_xml  # noqa: E402
from pptx.util import Inches  # noqa: E402

from src.extractors.pptx_extractor import PPTXExtractor  # noqa: E402

A_DECL = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'


def _set_bullet(paragraph, xml: str) -> None:
    """Give a paragraph an explicit bullet element (``<a:buNone/>`` etc.)."""
    element = xml.replace("/>", f" {A_DECL}/>", 1)
    paragraph._p.get_or_add_pPr().insert(0, pptx_parse_xml(element))


def _add_paragraphs(frame, items: list[tuple[str, int]]) -> list:
    paragraphs = []
    for index, (text, level) in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = text
        paragraph.level = level
        paragraphs.append(paragraph)
    return paragraphs


@pytest.fixture
def deck(tmp_path: Path) -> Path:
    prs = pptx.Presentation()

    # 1: body placeholder -- bulleted by default via the master's bodyStyle.
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "Agenda"
    body = s1.placeholders[1].text_frame
    paragraphs = _add_paragraphs(
        body, [("Intro", 0), ("Background", 1), ("Scope", 1), ("Plan", 0), ("No bullet here", 0)]
    )
    _set_bullet(paragraphs[-1], "<a:buNone/>")

    # 2: text boxes added out of reading order; numbering and bullets on text boxes.
    s2 = prs.slides.add_slide(prs.slide_layouts[6])
    s2.shapes.add_textbox(Inches(5), Inches(1), Inches(3), Inches(1)).text_frame.text = "Right column"
    left = s2.shapes.add_textbox(Inches(1), Inches(1.1), Inches(3), Inches(1)).text_frame
    left.text = "Left column with "
    run = left.paragraphs[0].add_run()
    run.text = "a link"
    run.hyperlink.address = "https://example.com/x"
    bottom = s2.shapes.add_textbox(Inches(1), Inches(4), Inches(6), Inches(2)).text_frame
    steps = _add_paragraphs(bottom, [("Step one", 0), ("Step two", 0), ("Detail", 1), ("Starred", 0), ("Plain line", 0)])
    _set_bullet(steps[0], '<a:buAutoNum type="arabicPeriod" startAt="3"/>')
    _set_bullet(steps[1], '<a:buAutoNum type="arabicPeriod" startAt="3"/>')
    _set_bullet(steps[2], '<a:buAutoNum type="alphaLcParenR"/>')
    _set_bullet(steps[3], '<a:buChar char="*"/>')

    # 3: a table with a merge, a chart, and speaker notes.
    s3 = prs.slides.add_slide(prs.slide_layouts[5])
    s3.shapes.title.text = "Numbers"
    table = s3.shapes.add_table(2, 2, Inches(1), Inches(1.5), Inches(4), Inches(1)).table
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(0, 0).text = "Region"
    table.cell(1, 0).text = "North"
    table.cell(1, 1).text = "a|b"
    data = CategoryChartData()
    data.categories = ["Q1", "Q2"]
    data.add_series("2024", (1.0, 2.5))
    data.add_series("2025", (3, 4))
    chart = s3.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(3.5), Inches(6), Inches(3), data
    ).chart
    chart.has_title = True
    chart.chart_title.text_frame.text = "Revenue"
    s3.notes_slide.notes_text_frame.text = "Remember the budget"

    # 4: a group (picture + caption) and a small logo with alt text above it.
    s4 = prs.slides.add_slide(prs.slide_layouts[6])
    group = s4.shapes.add_group_shape()
    photo = group.shapes.add_picture(io.BytesIO(_png(400, 300, "blue")), Inches(1), Inches(1))
    photo._element.nvPicPr.cNvPr.set("descr", "Team photo")
    group.shapes.add_textbox(Inches(1), Inches(5), Inches(3), Inches(1)).text_frame.text = "grouped caption"
    logo = s4.shapes.add_picture(io.BytesIO(_png(30, 30, "orange")), Inches(8), Inches(0.2))
    logo._element.nvPicPr.cNvPr.set("descr", "Logo")

    path = tmp_path / "fidelity.pptx"
    prs.save(path)
    return path


def test_pptx_bullets_numbering_order_tables_charts_notes(deck: Path):
    assert PPTXExtractor().extract(str(deck)).text == (
        "## Slide 1: Agenda\n\n"
        "- Intro\n"
        "  - Background\n"
        "  - Scope\n"
        "- Plan\n\n"
        "No bullet here\n\n"
        "## Slide 2\n\n"
        "Left column with [a link](https://example.com/x)\n\n"
        "Right column\n\n"
        "3. Step one\n"
        "4. Step two\n"
        "  a) Detail\n"
        "- Starred\n\n"
        "Plain line\n\n"
        "## Slide 3: Numbers\n\n"
        "| Region |  |\n"
        "| --- | --- |\n"
        "| North | a\\|b |\n\n"
        "**Chart: Revenue** (column clustered)\n\n"
        "| Category | 2024 | 2025 |\n"
        "| --- | --- | --- |\n"
        "| Q1 | 1 | 3 |\n"
        "| Q2 | 2.5 | 4 |\n\n"
        "### Notes\n\n"
        "Remember the budget\n\n"
        "## Slide 4\n\n"
        "![Logo](image:s4-1)\n\n"
        "![Team photo](image:s4-2)\n\n"
        "grouped caption"
    )


def test_pptx_image_ids_alt_and_group_positions(deck: Path):
    raw = PPTXExtractor().extract(str(deck))
    images = collect_images(raw.structure)
    assert [(img["id"], img["alt"], img["page"]) for img in images] == [("s4-2", "Team photo", 4)]
    # Group children are mapped through the group transform to slide points.
    assert images[0]["position"]["x0"] == pytest.approx(72)
    assert images[0]["position"]["y0"] == pytest.approx(72)
