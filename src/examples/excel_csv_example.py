"""
Example: Excel and CSV parsing with the parsers module.

Demonstrates:
- Excel file parsing (multi-sheet, merged cells, formulas)
- CSV file parsing (delimiter detection, encoding)
- Different chunking strategies
- HTML table conversion for LLMs
"""
import sys
import os
from pathlib import Path

# Add parent directory to path
parent_dir = Path(__file__).parent.parent
sys.path.insert(0, str(parent_dir))

from pipeline import DocumentPipeline


def create_sample_csv():
    """Create a sample CSV file for testing."""
    csv_content = """Product,Category,Price,Stock,Rating
Laptop,Electronics,999.99,50,4.5
Mouse,Electronics,24.99,200,4.2
Keyboard,Electronics,79.99,150,4.7
Monitor,Electronics,349.99,75,4.6
Desk,Furniture,299.99,30,4.3
Chair,Furniture,199.99,45,4.8
"""

    csv_file = Path("sample_data.csv")
    csv_file.write_text(csv_content)
    return csv_file


def create_sample_excel():
    """Create a sample Excel file with multiple sheets and features."""
    try:
        import openpyxl
        from openpyxl.styles import Font
    except ImportError:
        print("⚠️  openpyxl not installed. Run: pip install openpyxl")
        return None

    wb = openpyxl.Workbook()

    # Sheet 1: Sales Data
    ws1 = wb.active
    ws1.title = "Sales Data"

    # Add headers (bold)
    headers = ["Date", "Product", "Quantity", "Revenue"]
    for col, header in enumerate(headers, start=1):
        cell = ws1.cell(1, col, header)
        cell.font = Font(bold=True)

    # Add data
    data = [
        ["2024-01-15", "Laptop", 5, 4999.95],
        ["2024-01-16", "Mouse", 20, 499.80],
        ["2024-01-17", "Keyboard", 15, 1199.85],
        ["2024-01-18", "Monitor", 8, 2799.92],
    ]

    for row_idx, row_data in enumerate(data, start=2):
        for col_idx, value in enumerate(row_data, start=1):
            ws1.cell(row_idx, col_idx, value)

    # Merge cells example (total row)
    ws1.cell(6, 1, "TOTAL:")
    ws1.merge_cells('A6:C6')
    ws1.cell(6, 4, "=SUM(D2:D5)")  # Formula example

    # Sheet 2: Product Categories
    ws2 = wb.create_sheet("Categories")
    ws2.append(["Category", "Count", "Avg Price"])
    ws2.append(["Electronics", 4, 363.74])
    ws2.append(["Furniture", 2, 249.99])

    # Save
    excel_file = Path("sample_data.xlsx")
    wb.save(excel_file)
    return excel_file


def example_1_basic_csv_parsing():
    """Example 1: Basic CSV parsing with auto-detection."""
    print("\n" + "="*60)
    print("Example 1: Basic CSV Parsing")
    print("="*60)

    # Create sample CSV
    csv_file = create_sample_csv()

    try:
        # Create pipeline
        pipeline = DocumentPipeline()

        # Process CSV
        result = pipeline.process(str(csv_file))

        print(f"\n✓ Successfully parsed CSV file")
        print(f"  Format: {result['metadata']['format']}")
        print(f"  Delimiter: {result['metadata']['delimiter_name']}")
        print(f"  Encoding: {result['metadata']['encoding']}")
        print(f"  Rows: {result['metadata']['num_rows']}")
        print(f"  Columns: {result['metadata']['num_columns']}")
        print(f"  Total chunks: {result['metadata']['total_chunks']}")

        print(f"\n  First chunk preview:")
        first_chunk = result['chunks'][0]
        print(f"    Text (first 200 chars): {first_chunk['text'][:200]}...")
        print(f"    Metadata: {first_chunk['metadata']}")

    finally:
        # Cleanup
        csv_file.unlink(missing_ok=True)


def example_2_excel_multisheet_parsing():
    """Example 2: Excel with multiple sheets and formulas."""
    print("\n" + "="*60)
    print("Example 2: Excel Multi-Sheet Parsing")
    print("="*60)

    # Create sample Excel
    excel_file = create_sample_excel()

    if not excel_file:
        print("⚠️  Skipping example - openpyxl not installed")
        return

    try:
        # Create pipeline
        pipeline = DocumentPipeline()

        # Process Excel
        result = pipeline.process(str(excel_file))

        print(f"\n✓ Successfully parsed Excel file")
        print(f"  Format: {result['metadata']['format']}")
        print(f"  Total sheets: {result['metadata']['total_sheets']}")
        print(f"  Active sheets: {result['metadata']['active_sheets']}")
        print(f"  Sheet names: {', '.join(result['metadata']['sheet_names'])}")
        print(f"  Total chunks: {result['metadata']['total_chunks']}")

        # Show chunks
        for idx, chunk in enumerate(result['chunks']):
            print(f"\n  Chunk {idx}:")
            print(f"    Sheet: {chunk['metadata'].get('sheet_name', 'N/A')}")
            print(f"    Rows: {chunk['metadata'].get('num_rows', 0)}")
            print(f"    Columns: {chunk['metadata'].get('num_columns', 0)}")
            print(f"    Has formulas: {chunk['metadata'].get('has_formulas', False)}")
            print(f"    Has merged cells: {chunk['metadata'].get('has_merged_cells', False)}")
            print(f"    Text preview: {chunk['text'][:150]}...")

    finally:
        # Cleanup
        excel_file.unlink(missing_ok=True)


def example_3_excel_with_custom_chunking():
    """Example 3: Excel with different chunking strategies."""
    print("\n" + "="*60)
    print("Example 3: Excel Custom Chunking Strategies")
    print("="*60)

    excel_file = create_sample_excel()

    if not excel_file:
        print("⚠️  Skipping example - openpyxl not installed")
        return

    try:
        from chunking.strategies.excel_chunker import ExcelChunker
        from extractors.excel_extractor import ExcelExtractor

        # Extract
        extractor = ExcelExtractor()
        raw_doc = extractor.extract(str(excel_file))

        # Try different chunking strategies
        strategies = ["auto", "by_sheet", "by_table", "by_rows"]

        for strategy in strategies:
            chunker = ExcelChunker(strategy=strategy, chunk_size=50)
            chunks = chunker.chunk(raw_doc)

            print(f"\n  Strategy: {strategy}")
            print(f"    Total chunks: {len(chunks)}")

            for idx, chunk in enumerate(chunks[:2]):  # Show first 2 chunks
                print(f"    Chunk {idx}: {chunk.metadata.get('strategy', 'N/A')} - {len(chunk.text)} chars")

    finally:
        excel_file.unlink(missing_ok=True)


def example_4_html_table_conversion():
    """Example 4: Converting tables to HTML for LLMs."""
    print("\n" + "="*60)
    print("Example 4: HTML Table Conversion for LLMs")
    print("="*60)

    csv_file = create_sample_csv()

    try:
        from chunking.strategies.excel_chunker import ExcelChunker
        from extractors.csv_extractor import CSVExtractor

        # Extract CSV
        extractor = CSVExtractor()
        raw_doc = extractor.extract(str(csv_file))

        # Chunk with HTML conversion enabled
        chunker = ExcelChunker(convert_to_html=True)
        chunks = chunker.chunk(raw_doc)

        print(f"\n✓ Converted CSV to HTML table")
        print(f"  Total chunks: {len(chunks)}")
        print(f"\n  HTML output (first chunk):")
        print(chunks[0].text[:500])
        print("  ...")

    finally:
        csv_file.unlink(missing_ok=True)


def example_5_large_csv_chunking():
    """Example 5: Chunking large CSV files by rows."""
    print("\n" + "="*60)
    print("Example 5: Large CSV Chunking by Rows")
    print("="*60)

    # Create a larger CSV
    csv_content = "ID,Name,Value,Category\n"
    for i in range(200):
        csv_content += f"{i},Item_{i},{i*10},Category_{i%5}\n"

    csv_file = Path("large_data.csv")
    csv_file.write_text(csv_content)

    try:
        from chunking.strategies.excel_chunker import ExcelChunker
        from extractors.csv_extractor import CSVExtractor

        # Extract
        extractor = CSVExtractor()
        raw_doc = extractor.extract(str(csv_file))

        # Chunk by rows (50 rows per chunk)
        chunker = ExcelChunker(strategy="by_rows", chunk_size=50, preserve_headers=True)
        chunks = chunker.chunk(raw_doc)

        print(f"\n✓ Chunked large CSV (200 rows) into {len(chunks)} chunks")

        for idx, chunk in enumerate(chunks):
            row_start = chunk.metadata.get('row_start', 0)
            row_end = chunk.metadata.get('row_end', 0)
            print(f"  Chunk {idx}: Rows {row_start}-{row_end} ({row_end - row_start + 1} rows)")

        print(f"\n  First chunk preview:")
        print(f"  {chunks[0].text[:200]}...")

    finally:
        csv_file.unlink(missing_ok=True)


def main():
    """Run all examples."""
    print("\n" + "="*60)
    print("Excel/CSV Parsing Examples")
    print("="*60)

    print("\nℹ️  These examples demonstrate:")
    print("  - CSV parsing with auto-detection")
    print("  - Excel multi-sheet parsing")
    print("  - Merged cells and formula handling")
    print("  - Different chunking strategies")
    print("  - HTML table conversion for LLMs")
    print("  - Large file chunking")

    # Run examples
    try:
        example_1_basic_csv_parsing()
        example_2_excel_multisheet_parsing()
        example_3_excel_with_custom_chunking()
        example_4_html_table_conversion()
        example_5_large_csv_chunking()

        print("\n" + "="*60)
        print("✓ All examples completed successfully!")
        print("="*60 + "\n")

    except Exception as e:
        print(f"\n❌ Error running examples: {e}")
        import traceback
        traceback.print_exc()


if __name__ == "__main__":
    main()
